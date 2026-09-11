import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report_docx():
    doc = docx.Document()

    # Set A4 Page and Spiral Binding Margins (Extra margin on left punch edge)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)  # Reserved 1.25" for spiral binding punch edge
        section.right_margin = Inches(1.0)

    # Base typography
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(5)

    with open('RESEARCH_PROJECT_REPORT.md', 'r') as f:
        lines = f.readlines()

    table_lines = []
    in_code = False
    code_lines = []
    is_cover = True

    def style_table(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                tcPr = cell._tc.get_or_add_tcPr()
                if i == 0:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
                    tcPr.append(shd)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        for r in p.runs:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                            r.font.size = Pt(9.5)
                else:
                    if i % 2 == 1:
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                        tcPr.append(shd)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        for r in p.runs:
                            r.font.size = Pt(9.5)

    def process_table(t_lines):
        rows = [l.strip().strip('|').split('|') for l in t_lines if not '---' in l]
        if not rows: 
            return
        num_cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=num_cols)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                if c_idx < num_cols:
                    cell = t.cell(r_idx, c_idx)
                    cell.text = val.strip()
        style_table(t)
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code:
                in_code = False
                code_text = ''.join(code_lines)
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.2)
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(6)
                run = p_code.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith('|') and '|' in stripped[1:]:
            table_lines.append(line)
            i += 1
            continue
        else:
            if table_lines:
                process_table(table_lines)
                table_lines = []

        # Handle page breaks
        if stripped == '\\newpage':
            doc.add_page_break()
            is_cover = False
            i += 1
            continue

        if stripped == '---':
            i += 1
            continue

        # Handle markdown images ![caption](path)
        if stripped.startswith('![') and '](' in stripped and stripped.endswith(')'):
            img_path = stripped.split('](')[1][:-1].strip()
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Inches(5.6))
            i += 1
            continue

        # Handle figure captions *Figure X: ...*
        if stripped.startswith('*Figure ') and stripped.endswith('*'):
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            run_cap = p_cap.add_run(stripped[1:-1])
            run_cap.font.italic = True
            run_cap.font.size = Pt(9.5)
            run_cap.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Handle Headings
        if stripped.startswith('# '):
            heading_text = stripped[2:].strip()
            h = doc.add_heading(level=1)
            if is_cover:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if "XAVIER" in heading_text.upper() or "INSTITUTION" in heading_text.upper():
                    h.paragraph_format.space_before = Pt(30)
                    h.paragraph_format.space_after = Pt(6)
                    r = h.add_run(heading_text)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(17)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)
                elif "M.Sc" in heading_text:
                    h.paragraph_format.space_before = Pt(10)
                    h.paragraph_format.space_after = Pt(4)
                    r = h.add_run(heading_text)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(19)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    h.paragraph_format.space_before = Pt(25)
                    h.paragraph_format.space_after = Pt(15)
                    r = h.add_run(heading_text)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(21)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(30, 58, 138)
            else:
                h.paragraph_format.space_before = Pt(14)
                h.paragraph_format.space_after = Pt(5)
                r = h.add_run(heading_text)
                r.font.name = 'Calibri'
                r.font.size = Pt(18)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)

        elif stripped.startswith('## '):
            heading_text = stripped[3:].strip()
            h = doc.add_heading(level=2)
            if is_cover:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                h.paragraph_format.space_before = Pt(4)
                h.paragraph_format.space_after = Pt(4)
                r = h.add_run(heading_text)
                r.font.name = 'Calibri'
                r.font.size = Pt(18)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(4)
                r = h.add_run(heading_text)
                r.font.name = 'Calibri'
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)

        elif stripped.startswith('### '):
            heading_text = stripped[4:].strip()
            h = doc.add_heading(level=3)
            if is_cover:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                h.paragraph_format.space_before = Pt(2)
                h.paragraph_format.space_after = Pt(25)
                r = h.add_run(heading_text.upper())
                r.font.name = 'Calibri'
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(71, 85, 105)
            else:
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(3)
                r = h.add_run(heading_text)
                r.font.name = 'Calibri'
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = RGBColor(51, 65, 85)

        elif stripped.startswith('#### '):
            heading_text = stripped[5:].strip()
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(heading_text)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(71, 85, 105)

        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            text = stripped[2:].strip()
            parts = text.split('**')
            for idx, part in enumerate(parts):
                r = p.add_run(part)
                if idx % 2 == 1:
                    r.font.bold = True

        elif stripped:
            p = doc.add_paragraph()
            if is_cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.2
            else:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.15
            parts = stripped.split('**')
            for idx, part in enumerate(parts):
                r = p.add_run(part)
                if idx % 2 == 1:
                    r.font.bold = True

        i += 1

    if table_lines:
        process_table(table_lines)

    output_path = 'RESEARCH_PROJECT_REPORT.docx'
    doc.save(output_path)
    print(f'{output_path} generated successfully with centered cover and embedded figures!')

if __name__ == '__main__':
    create_report_docx()
